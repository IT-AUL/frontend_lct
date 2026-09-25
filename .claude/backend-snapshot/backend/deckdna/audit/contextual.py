"""Contextual VLM-аудит — 11 официальных вопросов к каждому слайду.

Реализует contextual-слой Render Arena (docs/pipeline/AUDIT_AND_REPAIR.md §3,
промпт ``prompts/contextual_audit/slide_checks.v1.yaml``): каждый слайд
рендерится в PNG (тот же soffice→pdftoppm путь, что у HTML-экспорта), VLM
получает картинку + извлечённый текст слайда + title_intent из DeckPlan +
evidence_excerpt из EvidenceGraph + язык колоды и отвечает JSON-verdict'ами
по 10 проверкам (вопрос 10 объединяет language consistency и logical flow —
мапится на ``deck.logical_flow``; ``content.language_consistency`` покрывается
этим же вердиктом и не заводится отдельно).

Маппинг вердиктов в ``AuditIssue`` (severity по frozen Appendix):

- ``fail`` + ``confidence >= confidence_threshold`` → issue
  (``deterministic=False``, verdict+confidence в evidence);
- ``uncertain`` → НЕ issue: честная неопределённость, не false positive.
  Понижение uncertain→warning (политика ``uncertain_downgrade_to`` из
  configs/audit.default.yaml) — осознанно отложено: выдавать warning за
  "не знаю" — это искусственный сигнал; disclosure реализуем, когда у
  issue появится поле для него;
- ``pass`` или ``fail`` ниже порога уверенности → не issue.

Работает против любого OpenAI-совместимого vision endpoint через
``gateway.vision_json`` — включая VK inference, когда его спецификация
будет опубликована. В текущем контуре VK endpoint'а нет; тесты идут
через MockProvider с фикстурными ответами.
"""

from __future__ import annotations

import asyncio
import base64
import tempfile
import uuid
from pathlib import Path
from typing import Literal

from pydantic import BaseModel

from deckdna.audit.issues import AuditIssue
from deckdna.contracts.deck_plan import DeckPlan
from deckdna.contracts.evidence_graph import EvidenceGraph
from deckdna.errors import DeckDNAError
from deckdna.pptx.exporting.render import render_slides_png
from deckdna.providers.base import ModelGateway

RULE_VERSION = "audit-contextual/0.1.0"
PROMPT_NAME = "slide_checks"
PROMPT_REF = "contextual_slide_audit@1.0.0"
DEFAULT_CONFIDENCE_THRESHOLD = 0.5

# Слайды в колоде независимы — гоняем VLM-вызовы конкурентно, а не по
# одному: с reasoning-моделью round-trip держится в районе 30-90с, и
# последовательный цикл на 12 слайдах давал ~10+ минут только на аудит
# одного варианта (замерено вживую на qwen3.8-27b/OpenRouter). Потолок
# ограничивает нагрузку на endpoint, не даёт бюджету расти — тот же
# набор вызовов, просто не в очередь.
DEFAULT_MAX_CONCURRENT_VLM_CALLS = 4

# Единичный отказ модели ("provider returned invalid structured output")
# не должен ронять весь ~10-минутный VLM-аудит варианта целиком — один
# повтор того же запроса. Другие коды ошибок (auth/5xx/schema-mismatch
# после повтора) остаются фатальными — они не транзиентный глюк ответа.
_RETRY_ON_ERROR_CODE = "structured_output_invalid"


async def _vision_json_with_retry(
    gateway: ModelGateway,
    prompt_name: str,
    images: list[str],
    payload: dict,
    schema: type[BaseModel],
    *,
    retries: int = 1,
) -> BaseModel:
    attempt = 0
    while True:
        try:
            return await gateway.vision_json(prompt_name, images, payload, schema)
        except DeckDNAError as exc:
            if exc.code != _RETRY_ON_ERROR_CODE or attempt >= retries:
                raise
            attempt += 1


class SlideCheckVerdict(BaseModel):
    """Один verdict из ответа VLM по slide_checks.v1.yaml."""

    check: str
    verdict: Literal["pass", "fail", "uncertain"]
    rationale: str | None = None
    confidence: float = 0.0


class SlideChecksResult(BaseModel):
    """Ответ VLM: verdict'ы по всем проверкам одного слайда."""

    verdicts: list[SlideCheckVerdict]


# check N (нумерация вопросов в slide_checks.v1.yaml) → frozen rule_code
# и severity из OFFICIAL_TRACEABILITY.md Appendix. Вопрос 10 объединяет
# language consistency и logical flow — мапим на deck.logical_flow.
_CHECK_RULES: dict[int, tuple[str, str]] = {
    1: ("content.conclusion_title", "warning"),
    2: ("content.title_alignment", "error"),
    3: ("content.single_message", "warning"),
    4: ("content.source_support", "blocker"),
    5: ("content.nonempty", "error"),
    6: ("content.visual_relevance", "warning"),
    7: ("content.prompt_leakage", "error"),
    8: ("content.spelling", "warning"),
    9: ("content.data_relevance", "warning"),
    10: ("deck.logical_flow", "warning"),
}


def _check_index(raw: str) -> int | None:
    """check приходит как "1"/"check_1"/"check 1" — извлекаем номер."""
    digits = "".join(c for c in raw if c.isdigit())
    return int(digits) if digits else None


def _slide_texts(pptx_path: Path) -> list[str]:
    """Текст каждого слайда (зеркало render._slide_texts, тот же код)."""
    from pptx import Presentation

    texts: list[str] = []
    for slide in Presentation(str(pptx_path)).slides:
        parts = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        texts.append("\n".join(parts))
    return texts


def _evidence_excerpt(
    slide_index: int,
    deck_plan: DeckPlan | None,
    graph: EvidenceGraph | None,
    *,
    max_chars: int = 2000,
) -> str:
    """Тексты evidence-узлов, релевантных слайду (по evidence_ids плана)."""
    if graph is None:
        return ""
    texts: dict[str, str] = {}
    for node in graph.nodes:
        if node.text:
            texts[node.id] = node.text
    ids: list[str] | None = None
    if deck_plan is not None:
        sp = _slide_plan_at(deck_plan, slide_index)
        if sp is not None:
            ids = sp.evidence_ids
    chosen = [texts[i] for i in ids if i in texts] if ids else list(texts.values())
    return "\n".join(chosen)[:max_chars]


def _png_data_url(png: Path) -> str:
    return "data:image/png;base64," + base64.b64encode(png.read_bytes()).decode()


async def run_contextual_audit(
    deck_path: str | Path,
    gateway: ModelGateway,
    *,
    deck_plan: DeckPlan | None = None,
    evidence_graph: EvidenceGraph | None = None,
    deck_language: str = "ru",
    confidence_threshold: float = DEFAULT_CONFIDENCE_THRESHOLD,
    first: int = 1,
    last: int | None = None,
    audit_run_id: str | None = None,
    deck_revision: int = 0,
    keep_images: str | Path | None = None,
    max_concurrent: int = DEFAULT_MAX_CONCURRENT_VLM_CALLS,
) -> list[AuditIssue]:
    """VLM-аудит каждого слайда колоды; возвращает contextual AuditIssue'ы.

    ``deck_plan``/``evidence_graph`` опциональны: без них title_intent и
    evidence_excerpt пустые — проверки на фактах/выводах честно вернут
    uncertain, а не выдуманный fail. ``keep_images`` оставляет PNG-рендеры.
    Слайды проверяются конкурентно (до ``max_concurrent`` одновременных
    vision_json-вызовов) — порядок итоговых issues и нумерация ``seq``
    остаются позиционными по слайду, не по порядку завершения запроса.
    """
    pptx_path = Path(deck_path)
    run_id = audit_run_id or f"audit-ctx-{uuid.uuid4().hex[:8]}"
    texts = _slide_texts(pptx_path)
    issues: list[AuditIssue] = []
    seq = 0

    with tempfile.TemporaryDirectory(prefix="deckdna-ctx-") as tmp:
        out_dir = Path(keep_images) if keep_images else Path(tmp)
        pngs = render_slides_png(pptx_path, out_dir, first=first, last=last)

        semaphore = asyncio.Semaphore(max(1, max_concurrent))

        async def _check_slide(pos: int, png: Path) -> tuple[int, SlideChecksResult]:
            slide_index = first + pos
            payload = {
                # input_fields промпта: slide_image — идентификатор рендера;
                # сама картинка идёт в сообщении как image_url (data URL).
                "slide_image": slide_index,
                "slide_text": texts[slide_index - 1] if slide_index - 1 < len(texts) else "",
                "title_intent": _title_intent(deck_plan, slide_index),
                "evidence_excerpt": _evidence_excerpt(
                    slide_index, deck_plan, evidence_graph
                ),
                "deck_language": deck_language,
            }
            async with semaphore:
                result = await _vision_json_with_retry(
                    gateway, PROMPT_NAME, [_png_data_url(png)], payload, SlideChecksResult
                )
            return slide_index, result

        per_slide = await asyncio.gather(
            *(_check_slide(pos, png) for pos, png in enumerate(pngs))
        )

        for slide_index, result in per_slide:
            for verdict in result.verdicts:
                if verdict.verdict != "fail":
                    continue
                if verdict.confidence < confidence_threshold:
                    continue
                rule = _CHECK_RULES.get(_check_index(verdict.check) or -1)
                if rule is None:
                    continue
                rule_code, severity = rule
                seq += 1
                issues.append(
                    AuditIssue(
                        id=f"{run_id}-{rule_code}-s{slide_index}-{seq}",
                        audit_run_id=run_id,
                        deck_revision=deck_revision,
                        rule_code=rule_code,
                        severity=severity,
                        message=(
                            verdict.rationale
                            or f"VLM: проверка {verdict.check} = fail"
                        ),
                        deterministic=False,
                        slide_index=slide_index,
                        measured_value=f"confidence {verdict.confidence:.2f}",
                        threshold=confidence_threshold,
                        confidence=verdict.confidence,
                        evidence=[
                            {
                                "kind": "vlm_verdict",
                                "ref": f"slide[{slide_index}]",
                                "detail": (
                                    f"{verdict.check}: fail "
                                    f"(confidence {verdict.confidence:.2f}); "
                                    f"{verdict.rationale or 'no rationale'}"
                                ),
                            }
                        ],
                        provenance={
                            "rule_version": RULE_VERSION,
                            "prompt": PROMPT_REF,
                        },
                    )
                )
    return issues


def _slide_plan_at(deck_plan: DeckPlan, slide_index: int):
    """SlidePlan для N-го слайда колоды — позиционно, не по sp.index.

    База sp.index не специфицирована контрактом: plan_deck выдаёт
    0-based, сторонний/LLM-план может быть 1-based — сверять index с
    1-based slide_index смещало бы граундинг на соседний слайд.
    """
    if not 1 <= slide_index <= len(deck_plan.slides):
        return None
    return deck_plan.slides[slide_index - 1]


def _title_intent(deck_plan: DeckPlan | None, slide_index: int) -> str:
    if deck_plan is None:
        return ""
    sp = _slide_plan_at(deck_plan, slide_index)
    return sp.title_intent if sp is not None else ""
